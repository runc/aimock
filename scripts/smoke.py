"""冒烟测试：按主流程顺序串起 server-v2 全部服务。

  a. GET  /api/sdm/tasks?user=李智勇          信息收集节点列任务
  a. GET  /api/sdm/tasks/{id}                 取任务基本信息
  b. POST /api/archive/parse                  上传示例压缩包并校验
  d. GET  /api/templates/by-task-name         按任务名取结构化模板
  e. GET  /api/standards/结构强度              取持久化标准
  c. POST /api/report/docx                    生成 Word 初版
  + 下载 docx 并用 python-docx 读回验证

用法：先启动服务（uv run uvicorn app.main:app --port 8200），
再执行 uv run python scripts/smoke.py [base_url]
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import httpx

BASE_DIR = Path(__file__).resolve().parent.parent
DEFAULT_BASE = "http://127.0.0.1:8200"
SAMPLE_ZIP = BASE_DIR / "data/sample/TL2A车架强度分析_二次处理_理想格式.zip"


def step(title: str) -> None:
    print(f"\n=== {title} ===")


def main() -> int:
    base = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_BASE
    if not SAMPLE_ZIP.exists():
        print(f"示例压缩包不存在：{SAMPLE_ZIP}\n先执行 uv run python scripts/make_sample_archive.py")
        return 1
    client = httpx.Client(base_url=base, timeout=60)

    step("0 health")
    health = client.get("/api/health").json()
    print(health)
    assert health["status"] == "ok"

    step("a1 列出当前用户待处理任务（信息收集节点）")
    tasks = client.get("/api/sdm/tasks", params={"user": "李智勇"}).json()
    for t in tasks["items"]:
        print(f"  {t['task_id']}  {t['task_name']}  类型={t['task_type']} pdca={t['pdca']} 状态={t['status']}")
    assert tasks["count"] >= 1
    target = next(t for t in tasks["items"] if t["task_type"] == "结构强度" and not t["pdca"])

    step("a2 取任务基本信息（HTTP 节点调 SDM）")
    detail = client.get(f"/api/sdm/tasks/{target['task_id']}").json()
    print(f"  任务名：{detail['task_name']}\n  任务单：{detail['task_no']}  报告编号：{detail['report_no']}"
          f"  项目号：{detail['project_no']}\n  签署：{detail['authors']}"
          f"\n  关联文档：{[d['name'] for d in detail['workspace_docs']]}")
    assert detail["task_no"] and detail["authors"]["编制"]

    step("b 上传二次处理压缩包并解析校验")
    with open(SAMPLE_ZIP, "rb") as f:
        parsed = client.post(
            "/api/archive/parse",
            params={"task_id": target["task_id"]},
            files={"file": (SAMPLE_ZIP.name, f, "application/zip")},
        ).json()
    print(f"  valid={parsed['valid']}  部件={parsed['summary']['parts']}  "
          f"工况={parsed['summary']['conditions']}  图片={parsed['summary']['images']}")
    print(f"  表格：{[(t['file'], t['kind']) for t in parsed['tables']]}")
    print(f"  工况清单：{parsed['conditions']}")
    print(f"  errors={parsed['errors']}\n  warnings={parsed['warnings']}")
    assert parsed["valid"], f"解析失败：{parsed['errors']}"
    assert parsed["summary"]["parts"] == 6
    wjgb = next(r for r in parsed["result_table"] if r["part"] == "外加强板")
    assert wjgb["stresses"]["吊点X向"] == 353.0

    step("d 按任务名取结构化模板")
    tpl = client.get("/api/templates/by-task-name", params={"task_name": detail["task_name"]}).json()
    print(f"  模板：{tpl['template_id']}（{tpl['name']}）  章节：{len(tpl['sections'])}")
    assert tpl["template_id"] == "cae_strength_report_v1"
    # 任务类型解析：PDCA 任务与不支持类型
    pdca_tpl = client.get("/api/templates/by-task-name",
                          params={"task_name": "D320随车吊TL2A车架强度分析（PDCA整改复核）"}).json()
    assert pdca_tpl["task_type_info"]["pdca"] is True
    r = client.get("/api/templates/by-task-name", params={"task_name": "X23驾驶室正面碰撞分析"})
    assert r.status_code == 404, "碰撞任务应无模板（404）"

    step("e 取按任务类型持久化的标准二维表")
    std = client.get("/api/standards/结构强度").json()
    print(f"  {std['name']} v{std['version']}：{len(std['condition_factors'])} 工况系数，"
          f"{len(std['material_yields'])} 材料屈服强度")

    step("c 生成 Word 初版")
    gen = client.post("/api/report/docx", json={
        "task": detail,
        "result": parsed,
    }).json()
    print(f"  file_url={gen['file_url']}")
    print(f"  判定摘要：{gen['judge_summary']}")
    fails = [v for v in gen["verdicts"]["rows"] if v["verdict"] == "不合格"]
    print(f"  不合格 {len(fails)} 项：{[(f['part'], f['condition'], f['stress_mpa'], f['limit_mpa']) for f in fails]}")
    print(f"  warnings：{gen['warnings']}")
    assert "外加强板" in gen["judge_summary"] and "不合格" in gen["judge_summary"]

    step("下载 docx 并读回验证")
    resp = client.get(gen["file_url"].replace(base, ""))
    assert resp.status_code == 200, f"下载失败 {resp.status_code}"
    from docx import Document

    doc = Document(io.BytesIO(resp.content))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    images = sum(1 for s in doc.inline_shapes)
    print(f"  段落 {len(texts)}  表格 {len(doc.tables)}  图片 {images}")
    print(f"  首段：{texts[0]}")
    conclusion = next(t for t in texts if t.startswith("共 "))
    print(f"  结论段：{conclusion[:80]}...")
    assert len(doc.tables) >= 6 and images >= 3

    print("\n✅ 冒烟通过：a/b/c/d + 标准持久化 全链路 OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

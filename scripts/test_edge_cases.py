"""边界用例验证（不进冒烟主链路）：分开格式三表合并 / PDCA 任务出稿 / 错误路径。

用法：服务启动后 uv run python scripts/test_edge_cases.py
"""

from __future__ import annotations

import io
import sys
import zipfile

import httpx

BASE = "http://127.0.0.1:8200"


def sep_zip() -> bytes:
    """分开格式：材料表（无屈服）+ 标准应力表 + 非标应力表，行名用「零件」。"""
    material = "零件,材料\n纵梁,DL590\n纵梁加强板,DL591\n外加强板,DL592\n"
    stress = ("零件 / 工况,弯曲工况 bend,扭转工况 torsion1,转弯工况 Cornering,"
              "扭转+转弯工况 tor+cor,制动工况 brake,双扭工况 bad\n"
              "纵梁,118.3,225.7,130.5,217.1,191.1,325.1\n"
              "纵梁加强板,80.6,158.2,101.1,146.5,93.4,231.8\n"
              "外加强板,76.5,151.9,92.9,173.2,65.7,200.8\n")
    nonstd = ("部件/零件 / 工况,吊点X向工况,吊点XY中间工况,吊点Y向工况\n"
              "纵梁,241.7,187.6,77.9\n纵梁加强板,302.8,200.0,66.2\n外加强板,353.0,242.0,105.9\n")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("材料表.csv", material.encode("gbk"))
        zf.writestr("应力表.csv", stress.encode("utf-8-sig"))
        zf.writestr("非标应力表.csv", nonstd.encode("utf-8-sig"))
    return buf.getvalue()


def main() -> int:
    client = httpx.Client(base_url=BASE, timeout=60)
    ok = True

    # 1. 分开格式三表合并（gbk 材料表 + utf-8 应力表）
    r = client.post("/api/archive/parse", files={"file": ("分开格式.zip", sep_zip(), "application/zip")}).json()
    print("【分开格式】valid=", r["valid"], " tables=", [(t["file"], t["kind"]) for t in r["tables"]],
          " parts=", [p["part"] for p in r["result_table"]])
    zl = next(p for p in r["result_table"] if p["part"] == "纵梁")
    print("  纵梁：材料=", zl["material"], " 工况数=", len(zl["stresses"]),
          " 吊点X向=", zl["stresses"].get("吊点X向"), " 待确认告警=", len(r["warnings"]))
    assert r["valid"] and zl["material"] == "DL590" and zl["stresses"].get("吊点X向") == 241.7
    assert len(zl["stresses"]) == 9

    # 2. 分开格式出稿：标准工况用持久化企标系数判定，非标工况待确认
    task = client.get("/api/sdm/tasks/SDM-SIM-2026-1187").json()
    gen = client.post("/api/report/docx", json={"task": task, "result": r}).json()
    print("【分开格式出稿】", gen["judge_summary"], " 待确认=", gen["verdicts"]["pending_count"])
    assert gen["verdicts"]["pending_count"] == 9  # 3 部件 × 3 非标工况全部待确认
    assert "合格" in gen["verdicts"]["judge_summary"] or "待确认" in gen["verdicts"]["judge_summary"]

    # 3. PDCA 任务出稿
    pdca_task = client.get("/api/sdm/tasks/SDM-SIM-2026-1201").json()
    assert pdca_task["pdca"] is True
    ideal = open("data/sample/TL2A车架强度分析_二次处理_理想格式.zip", "rb").read()
    pr = client.post("/api/archive/parse", files={"file": ("理想.zip", ideal, "application/zip")}).json()
    gen2 = client.post("/api/report/docx", json={
        "task": pdca_task, "result": pr, "file_name": "PDCA复核报告.docx"}).json()
    print("【PDCA 出稿】", gen2["file_name"], "→", gen2["judge_summary"])
    from docx import Document

    doc = Document(io.BytesIO(client.get(gen2["file_url"].replace(BASE, "")).content))
    all_texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        all_texts += [c.text for row in table.rows for c in row.cells]
    purpose = next(t for t in all_texts if t.startswith("本任务为 PDCA"))
    print("  分析目的：", purpose[:50], "...")
    assert purpose

    # 4. 错误路径：非压缩包 → 415；损坏 zip → valid=false；缺材料 → error
    r415 = client.post("/api/archive/parse", files={"file": ("a.txt", b"hello", "text/plain")})
    print("【非压缩包】status=", r415.status_code)
    assert r415.status_code == 415
    bad = zipfile.ZipFile(io.BytesIO(), "w")
    bad_fp = io.BytesIO()
    with zipfile.ZipFile(bad_fp, "w") as z:
        z.writestr("x.csv", "部件,弯曲工况\n")
    raw = bad_fp.getvalue()[:-5]  # 截断制造损坏
    rbad = client.post("/api/archive/parse", files={"file": ("bad.zip", raw, "application/zip")})
    print("【损坏 zip】status=", rbad.status_code, " valid=", rbad.json().get("valid"))
    assert rbad.json().get("valid") is False

    no_material = "部件,弯曲工况\n纵梁,118.3\n"
    nm_buf = io.BytesIO()
    with zipfile.ZipFile(nm_buf, "w") as z:
        z.writestr("t.csv", no_material.encode("utf-8"))
    rnm = client.post("/api/archive/parse", files={"file": ("nm.zip", nm_buf.getvalue(), "application/zip")}).json()
    print("【缺材料】valid=", rnm["valid"], " errors=", [e["code"] for e in rnm["errors"]])
    assert rnm["valid"] is False and any(e["code"] == "MATERIAL_MISSING" for e in rnm["errors"])

    # 5. 碰撞任务无模板 → 404
    r404 = client.post("/api/report/docx", json={"task": {"task_name": "X23驾驶室正面碰撞分析"}})
    print("【碰撞任务出稿】status=", r404.status_code)
    assert r404.status_code == 404

    print("\n✅ 边界用例全部通过")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
